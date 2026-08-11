from pathlib import Path

from docx import Document as DocxDocument
from reportlab.pdfgen import canvas

from app.ingestion.parser import parse_document


def test_parse_txt():
    path = Path("data/sample/test.txt")

    pages = parse_document(str(path))

    assert len(pages) == 1
    assert "Acme Supplies" in pages[0]["text"]
    assert pages[0]["page_number"] is None


def test_parse_pdf(tmp_path):
    pdf_path = tmp_path / "test.pdf"

    pdf = canvas.Canvas(str(pdf_path))
    pdf.drawString(100, 750, "Vendor: Acme Supplies")
    pdf.drawString(100, 730, "Contract Value: $50,000")
    pdf.save()

    pages = parse_document(str(pdf_path))

    assert len(pages) == 1
    assert "Acme Supplies" in pages[0]["text"]
    assert pages[0]["page_number"] == 1


def test_parse_docx(tmp_path):
    docx_path = tmp_path / "test.docx"

    document = DocxDocument()
    document.add_paragraph("Vendor: Acme Supplies")
    document.add_paragraph("Contract Value: $50,000")
    document.save(str(docx_path))

    pages = parse_document(str(docx_path))

    assert len(pages) == 1
    assert "Acme Supplies" in pages[0]["text"]